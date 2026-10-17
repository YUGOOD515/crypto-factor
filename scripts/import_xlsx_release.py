#!/usr/bin/env python3
"""Import a quarterly factor bundle (.xlsx) into data/releases/<RELEASE_ID>/.

Usage:
    python3 scripts/import_xlsx_release.py path/to/bundle.xlsx 2026Q2 [--latest]

The workbook is expected to follow the delivered bundle layout:
one long sheet per factor (master_id, signal_id, factor_name, formation_date,
return_date, factor_return) plus ALL_FACTORS_WIDE. Each release is written to
its own folder, so earlier releases stay in the repository (history is kept),
while the website only publishes the release marked as "latest" in
data/releases/manifest.json.

Requires: openpyxl (pip install openpyxl)
"""
import csv
import datetime as dt
import json
import pathlib
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
RELEASES = ROOT / "data" / "releases"

# wide-table code -> sheet name in the delivered workbook
SHEETS = {
    "MKT": "MKT",
    "SMB": "SMB",
    "MOM": "MOM",
    "MOM3w": "MOM3w",
    "M007_GLASSNODE": "M007_GLASSNODE",
    "M007_DUNE": "M007_DUNE",
    "CANCHOR": "CANCHOR_CLEAN",
    "ST": "ST",
}
LONG_HEADER = ["master_id", "signal_id", "factor_name", "formation_date", "return_date", "factor_return"]


def fmt_date(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10]


def fmt_ret(v):
    if v is None or v == "":
        return ""
    return f"{float(v):.4f}"


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    src = pathlib.Path(sys.argv[1])
    release_id = sys.argv[2]
    make_latest = "--latest" in sys.argv

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    out = RELEASES / release_id / "factors"
    out.mkdir(parents=True, exist_ok=True)

    coverage = {}
    for code, sheet in SHEETS.items():
        rows = list(wb[sheet].iter_rows(values_only=True))
        header = [str(h) for h in rows[0]]
        assert header == LONG_HEADER, f"{sheet}: unexpected header {header}"
        body = [r for r in rows[1:] if r and r[0] is not None]
        with open(out / f"{code}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh, lineterminator="\n")
            w.writerow(LONG_HEADER)
            for r in body:
                w.writerow([r[0], r[1], r[2], fmt_date(r[3]), fmt_date(r[4]), fmt_ret(r[5])])
        valid = [r for r in body if r[5] is not None]
        coverage[code] = {
            "master_id": body[0][0],
            "signal_id": body[0][1],
            "factor_name": body[0][2],
            "n_rows": len(body),
            "n_nonmissing": len(valid),
            "formation_start": fmt_date(valid[0][3]),
            "formation_end": fmt_date(valid[-1][3]),
            "return_end": fmt_date(valid[-1][4]),
        }

    rows = list(wb["ALL_FACTORS_WIDE"].iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    with open(out / "ALL_FACTORS_WIDE.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        for r in rows[1:]:
            if not r or r[0] is None:
                continue
            w.writerow([fmt_date(r[0]), fmt_date(r[1])] + [fmt_ret(v) for v in r[2:]])

    release_meta = {
        "id": release_id,
        "source_workbook": src.name,
        "frequency": "weekly",
        "data_through": max(c["return_end"] for c in coverage.values()),
        "factors": coverage,
    }
    (RELEASES / release_id / "release.json").write_text(json.dumps(release_meta, indent=2) + "\n", encoding="utf-8")

    manifest_path = RELEASES / "manifest.json"
    manifest = json.loads(manifest_path.read_text()) if manifest_path.exists() else {"latest": None, "releases": []}
    if release_id not in [r["id"] for r in manifest["releases"]]:
        manifest["releases"].append({"id": release_id, "data_through": release_meta["data_through"]})
    if make_latest or manifest["latest"] is None:
        manifest["latest"] = release_id
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(coverage)} factors into {out} (latest = {manifest['latest']})")


if __name__ == "__main__":
    main()
